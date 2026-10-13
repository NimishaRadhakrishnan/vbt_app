"""Dealer "stock left" report, laid out like the company's stock sheet.

One block per dealer. Inside a block: Category | Product | Properties |
one column per pack size (50ml ... 10Kg) | Remarks, with the quantity left
in the dealer's shop under the matching pack size.

Pack sizes are not stored separately: every pack is its own product named
like "V-NITRO (Liquid 5 Lit)". This module splits that name the same way
the web/mobile stock sheet does, then files the pack under one of the
standard columns. A pack that fits none of them still gets its own extra
column at the end, so no quantity is ever dropped from the report.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

# Column order of the company's sheet.
STANDARD_PACKS = [
    "50ml", "100ml", "250ml", "500ml", "950ml", "1L", "5L", "10L",
    "250g", "500g", "900g", "1Kg", "1.5Kg", "5Kg", "10Kg",
]
_PACK_HEADINGS = {"1.5Kg": "1.5 KG"}  # the sheet spells this one differently

_NAME_RE = re.compile(r"^(.*?)\s*\((.*)\)\s*$")
_PROPS_RE = re.compile(r"^([A-Za-z]+)\s+(\d.*)$")
_PACK_RE = re.compile(r"^(\d+(?:\.\d+)?)\s*([A-Za-z]+)$")


def split_product_name(name: str) -> tuple[str, str, str]:
    """'V-NITRO (Liquid 5 Lit)' -> ('V-NITRO', 'Liquid', '5 Lit')."""
    m = _NAME_RE.match(name or "")
    base = ((m.group(1) if m else name) or "").strip()
    variant = (m.group(2) if m else "").strip()
    v = _PROPS_RE.match(variant)
    if v:
        return base, v.group(1), v.group(2).strip()
    return base, "Standard", variant or "Single pack"


def normalise_pack(pack: str) -> str:
    """'5 Lit' -> '5L', '500 ML' -> '500ml', '1 Kg' -> '1Kg'. Unknown text
    is returned trimmed, which becomes an extra column."""
    m = _PACK_RE.match(re.sub(r"\s+", "", pack or ""))
    if not m:
        return (pack or "").strip()
    number, unit = m.group(1), m.group(2).lower()
    if number.endswith(".0"):
        number = number[:-2]
    if unit in ("ml", "milliliter", "millilitre"):
        return f"{number}ml"
    if unit in ("l", "lit", "ltr", "litre", "liter", "ltrs", "lits"):
        return f"{number}L"
    if unit in ("g", "gm", "gms", "gram", "grams"):
        return f"{number}g"
    if unit in ("kg", "kgs", "kilo"):
        return f"{number}Kg"
    return (pack or "").strip()


@dataclass
class StockRow:
    dealer: str
    district: Optional[str]
    officer: Optional[str]
    visit_date: str
    category: str
    product_name: str
    quantity_left: int
    remarks: Optional[str] = None


@dataclass
class _Line:
    category: str
    product: str
    properties: str
    quantities: dict[str, int] = field(default_factory=dict)
    remarks: list[str] = field(default_factory=list)


@dataclass
class DealerBlock:
    dealer: str
    district: Optional[str]
    officer: Optional[str]
    visit_date: str
    lines: list[_Line]


def build_blocks(rows: list[StockRow]) -> tuple[list[DealerBlock], list[str]]:
    """Group rows into dealer blocks and work out the pack columns.

    Returns (blocks, pack_columns). The columns are the standard ones in
    sheet order, followed by any
    unrecognised pack labels. All standard columns are always shown so the
    printout matches the company sheet.
    """
    by_dealer: dict[tuple[str, str], DealerBlock] = {}
    used: set[str] = set()
    for r in rows:
        key = (r.dealer, r.visit_date)
        block = by_dealer.setdefault(
            key, DealerBlock(r.dealer, r.district, r.officer, r.visit_date, [])
        )
        base, props, pack_raw = split_product_name(r.product_name)
        pack = normalise_pack(pack_raw)
        used.add(pack)
        line = next(
            (l for l in block.lines
             if l.category == r.category and l.product == base and l.properties == props),
            None,
        )
        if line is None:
            line = _Line(r.category or "Other", base, props)
            block.lines.append(line)
        line.quantities[pack] = line.quantities.get(pack, 0) + r.quantity_left
        if r.remarks:
            line.remarks.append(r.remarks)

    columns = list(STANDARD_PACKS) + sorted(used - set(STANDARD_PACKS))
    blocks = sorted(by_dealer.values(), key=lambda b: (b.dealer.lower(), b.visit_date))
    for b in blocks:
        b.lines.sort(key=lambda l: (l.category, l.product.lower(), l.properties))
    return blocks, columns


def _heading(pack: str) -> str:
    return _PACK_HEADINGS.get(pack, pack)


GREEN = "3F6B4F"


def render_xlsx(rows: list[StockRow], title: str) -> bytes:
    blocks, columns = build_blocks(rows)
    wb = Workbook()
    ws = wb.active
    ws.title = "Stock left"
    header_fill = PatternFill("solid", fgColor=GREEN)
    block_fill = PatternFill("solid", fgColor="E7F0EA")
    thin = Side(style="thin", color="BBBBBB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    last_col = 3 + len(columns) + 1

    ws.cell(row=1, column=1, value=title).font = Font(bold=True, size=14)
    row_no = 3
    if not blocks:
        ws.cell(row=row_no, column=1, value="No stock sheet entries for this selection.")
    for b in blocks:
        info = f"{b.dealer}" + (f" - {b.district}" if b.district else "")
        info += f"   |   Officer: {b.officer or '-'}   |   Visit date: {b.visit_date}"
        c = ws.cell(row=row_no, column=1, value=info)
        c.font = Font(bold=True)
        for col in range(1, last_col + 1):
            ws.cell(row=row_no, column=col).fill = block_fill
        row_no += 1
        heads = ["Category", "Product", "Properties"] + [_heading(p) for p in columns] + ["Remarks"]
        for i, h in enumerate(heads, start=1):
            cell = ws.cell(row=row_no, column=i, value=h)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border
        row_no += 1
        for line in b.lines:
            values = [line.category, line.product, line.properties]
            values += [line.quantities.get(p) for p in columns]
            values.append("; ".join(line.remarks))
            for i, v in enumerate(values, start=1):
                cell = ws.cell(row=row_no, column=i, value=v)
                cell.border = border
                if 4 <= i < last_col:
                    cell.alignment = Alignment(horizontal="center")
            row_no += 1
        row_no += 1  # blank row between dealers

    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 14
    for i in range(4, last_col):
        ws.column_dimensions[get_column_letter(i)].width = 9
    ws.column_dimensions[get_column_letter(last_col)].width = 28
    ws.freeze_panes = "A3"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def render_pdf(rows: list[StockRow], title: str) -> bytes:
    blocks, columns = build_blocks(rows)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=landscape(A4), leftMargin=24, rightMargin=24, topMargin=24, bottomMargin=24,
        title=title,
    )
    styles = getSampleStyleSheet()
    cell = ParagraphStyle("cell", parent=styles["Normal"], fontSize=7, leading=8.5)
    head = ParagraphStyle("head", parent=cell, textColor=colors.white, fontName="Helvetica-Bold", alignment=1)
    story: list = [Paragraph(title, styles["Title"]), Spacer(1, 6)]
    if not blocks:
        story.append(Paragraph("No stock sheet entries for this selection.", styles["Normal"]))

    page_width = landscape(A4)[0] - 48
    fixed = [70, 100, 50]
    remarks_w = 90
    pack_w = max(22, (page_width - sum(fixed) - remarks_w) / max(1, len(columns)))
    col_widths = fixed + [pack_w] * len(columns) + [remarks_w]

    for b in blocks:
        info = f"<b>{b.dealer}</b>" + (f" - {b.district}" if b.district else "")
        info += f" &nbsp;|&nbsp; Officer: {b.officer or '-'} &nbsp;|&nbsp; Visit date: {b.visit_date}"
        story.append(Paragraph(info, styles["Normal"]))
        story.append(Spacer(1, 3))
        data = [[Paragraph(h, head) for h in
                 ["Category", "Product", "Properties"] + [_heading(p) for p in columns] + ["Remarks"]]]
        for line in b.lines:
            data.append(
                [Paragraph(line.category, cell), Paragraph(line.product, cell), Paragraph(line.properties, cell)]
                + [Paragraph(str(line.quantities[p]) if p in line.quantities else "", cell) for p in columns]
                + [Paragraph("; ".join(line.remarks), cell)]
            )
        table = Table(data, colWidths=col_widths, repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + GREEN)),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#BBBBBB")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("ALIGN", (3, 1), (-2, -1), "CENTER"),
        ]))
        story.append(table)
        story.append(Spacer(1, 12))
    doc.build(story)
    return buf.getvalue()
